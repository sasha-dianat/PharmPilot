"""
Tests for the owner-fed reference ingestion engine (replaces TTAC reliance):
DOCX parsing (incl. Persian + tables + zero-dep XML fallback).
"""
import docx
import pytest

from services.ai.knowledge_engine.parsers.document_parser import DOCXParser


@pytest.fixture
def persian_docx(tmp_path):
    p = tmp_path / "metformin_fa.docx"
    d = docx.Document()
    d.add_heading("مونوگراف متفورمین", level=1)
    d.add_paragraph("متفورمین در بیماران با نارسایی کلیه (eGFR کمتر از ۳۰) منع مصرف دارد.")
    d.add_heading("عوارض جانبی:", level=2)
    d.add_paragraph("اسیدوز لاکتیک، تهوع، اسهال.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "دوز"
    t.rows[0].cells[1].text = "۵۰۰ میلی‌گرم"
    d.save(str(p))
    return str(p)


def test_docx_parses_persian_content(persian_docx):
    chunks = DOCXParser().parse(persian_docx)
    assert len(chunks) >= 1
    joined = " ".join(c.content for c in chunks)
    assert "متفورمین" in joined
    assert "eGFR" in joined  # mixed Persian + Latin preserved


def test_docx_captures_table_cells(persian_docx):
    chunks = DOCXParser().parse(persian_docx)
    joined = " ".join(c.content for c in chunks)
    assert "۵۰۰" in joined  # table cell value present


def test_docx_short_monograph_not_dropped(tmp_path):
    """Short reference paragraphs must survive the chunker's length floor."""
    p = tmp_path / "short.docx"
    d = docx.Document()
    d.add_paragraph("آسپرین در کودکان زیر ۱۲ سال منع مصرف دارد.")  # < 100 chars
    d.save(str(p))
    chunks = DOCXParser().parse(str(p))
    assert len(chunks) == 1
    assert "آسپرین" in chunks[0].content


def test_docx_xml_fallback(tmp_path, monkeypatch):
    """If python-docx is unavailable, the raw word/document.xml fallback works."""
    p = tmp_path / "fallback.docx"
    d = docx.Document()
    d.add_paragraph("این یک متن آزمایشی برای مسیر جایگزین است که باید استخراج شود.")
    d.save(str(p))

    parser = DOCXParser()
    # Force the python-docx path to fail so the XML fallback runs.
    import builtins
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("simulated missing python-docx")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    paras = parser._extract_paragraphs(str(p))
    assert any("آزمایشی" in pp for pp in paras)
